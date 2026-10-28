# -*- coding: utf-8 -*-
"""
docx_doc.py — Word(.docx) 보고서 읽기 / 서식 유지 수정 / 절·문단·표·사진 추가

원칙: 새로 넣는 모든 것은 '같은 종류의 기존 요소'를 복사해서 글만 바꾼다.
      → 글씨체(rPr)·문단 모양(pPr)·표 테두리(tblPr/tcPr)가 원래 보고서와 같게 유지된다.

필요: pip install python-docx
"""
import copy
import difflib
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Emu

from hwp_reader import HwpDoc, Table, heading_of

W_P, W_TBL, W_R, W_T = qn("w:p"), qn("w:tbl"), qn("w:r"), qn("w:t")
W_TAB, W_BR, W_CR = qn("w:tab"), qn("w:br"), qn("w:cr")
W_DRAWING, W_PICT, W_OBJECT = qn("w:drawing"), qn("w:pict"), qn("w:object")
W_PPR, W_RPR, W_SDT, W_SDTC = qn("w:pPr"), qn("w:rPr"), qn("w:sdt"), qn("w:sdtContent")
W_TR, W_TC, W_TCPR = qn("w:tr"), qn("w:tc"), qn("w:tcPr")
W_DELTEXT, W_INSTR = qn("w:delText"), qn("w:instrText")
W_DEL = qn("w:del")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

RE_CAP_ANY = re.compile(r"^\s*[<\[〈(]?\s*(그림|표|사진)\s*(\d+)\s*[.\-]\s*(\d+)")


# ====================================================================== 읽기

def _walk_tokens(p_el):
    """문단 안의 글자 조각을 문서 순서대로: ('t', w:t) | ('x', 요소, 글자) | ('pic', 요소)"""
    out = []

    def rec(el):
        for ch in el:
            tag = ch.tag
            if tag == W_T:
                out.append(("t", ch))
            elif tag == W_TAB:
                out.append(("x", ch, "\t"))
            elif tag in (W_BR, W_CR):
                out.append(("x", ch, "\n"))
            elif tag in (W_DRAWING, W_PICT, W_OBJECT):
                out.append(("pic", ch))           # 그림/글상자 속은 들어가지 않음
            elif tag in (W_PPR, W_RPR, W_DEL, W_DELTEXT, W_INSTR):
                continue
            elif tag == W_P or tag == W_TBL:
                continue
            else:
                rec(ch)
    rec(p_el)
    return out


def para_text(p_el):
    """(text, breaks, has_pic)"""
    buf, breaks, pic = [], [], False
    for tok in _walk_tokens(p_el):
        if tok[0] == "t":
            buf.append(tok[1].text or "")
        elif tok[0] == "x":
            breaks.append(sum(map(len, buf)))
            buf.append(tok[2])
        else:
            pic = True
            breaks.append(sum(map(len, buf)))
    return "".join(buf), breaks, pic


def read_docx(path, key=None, document=None):
    """document 를 주면 파일 대신 메모리의 문서를 다시 분석 (구조 추가 후)"""
    dx = document or Document(path)
    doc = HwpDoc(path, key or path)
    doc.fmt = "docx"
    doc.docx = dx
    body = dx.element.body
    _parse_container(doc, body, doc.blocks)
    return doc


def _parse_container(doc, el, sink):
    for ch in el:
        if ch.tag == W_P:
            _add_para(doc, ch, sink, "body")
        elif ch.tag == W_TBL:
            sink.append(("t", _parse_table(doc, ch)))
        elif ch.tag == W_SDT:
            c = ch.find(W_SDTC)
            if c is not None:
                _parse_container(doc, c, sink)


def _add_para(doc, p_el, sink, kind, table=None, row=-1, col=-1):
    text, breaks, pic = para_text(p_el)
    p = doc.new_para(text, breaks, kind=kind, table=table, row=row, col=col)
    p.ref = p_el
    p.pic = pic
    if sink is not None:
        sink.append(("p", p))
    return p


def _parse_table(doc, tbl_el):
    t = Table(len(doc.tables), 0, 0)
    t.ref = tbl_el
    doc.tables.append(t)
    vmerge_origin = {}                 # 열 → 세로 병합 시작 셀
    r = -1
    ncols = 0
    for tr in tbl_el.iter(W_TR):
        if tr.getparent() is not tbl_el:
            continue                    # 중첩 표의 행은 건너뜀
        r += 1
        c = 0
        for tc in tr.findall(W_TC):
            tcpr = tc.find(W_TCPR)
            span, vm = 1, None
            if tcpr is not None:
                gs = tcpr.find(qn("w:gridSpan"))
                if gs is not None:
                    span = int(gs.get(qn("w:val"), "1"))
                v = tcpr.find(qn("w:vMerge"))
                if v is not None:
                    vm = v.get(qn("w:val"), "continue")
            if vm == "continue" and c in vmerge_origin:
                vmerge_origin[c]["rs"] += 1
                c += span
                continue
            cell = t.add_cell(r, c, 1, span)
            cell["ref"] = tc
            if tcpr is not None and tcpr.find(qn("w:tcW")) is not None:
                try:
                    cell["w"] = int(tcpr.find(qn("w:tcW")).get(qn("w:w"), "0"))
                except ValueError:
                    pass
            if vm == "restart":
                vmerge_origin[c] = cell
            else:
                vmerge_origin.pop(c, None)
            for ch in tc:
                if ch.tag == W_P:
                    cell["paras"].append(_add_para(doc, ch, None, "cell", t, r, c))
                elif ch.tag == W_TBL:               # 중첩 표 → 셀 문단으로 평탄화
                    sub = _parse_table(doc, ch)
                    for key in sorted(sub.cells):
                        cell["paras"].extend(sub.cells[key]["paras"])
            c += span
        ncols = max(ncols, c)
    t.rows, t.cols = r + 1, ncols
    return t


# ====================================================================== 글자 수정 (서식 유지)

def _new_t(text):
    t = copy.deepcopy(_T_PROTO)
    t.text = text
    t.set(XML_SPACE, "preserve")
    return t


_T_PROTO = Document().element.body.makeelement(W_T, {})


def _mk(tag):
    return _T_PROTO.makeelement(tag, {})


def _pieces_to_elems(s):
    """'가\t나\n다' → [w:t, w:tab, w:t, w:br, w:t]"""
    out = []
    buf = ""
    for ch in s:
        if ch in "\t\n":
            if buf:
                out.append(_new_t(buf))
                buf = ""
            out.append(_mk(W_TAB if ch == "\t" else W_BR))
        else:
            buf += ch
    if buf:
        out.append(_new_t(buf))
    return out


def set_para_text(p_el, old, new):
    """문단 글자를 old → new 로. 바뀐 부분만 고치고, 새 글자는 바로 앞 글자의 서식(run)을 따른다."""
    toks = _walk_tokens(p_el)
    cur = "".join(tok[1].text or "" if tok[0] == "t" else (tok[2] if tok[0] == "x" else "") for tok in toks)
    if cur != old:
        raise ValueError("문단 내용이 예상과 다릅니다 (다른 곳에서 바뀜)")
    if old == new:
        return
    if not any(tok[0] == "t" for tok in toks):
        _ensure_text_run(p_el)
        toks = _walk_tokens(p_el)
    sm = difflib.SequenceMatcher(None, old, new, autojunk=False)
    ops = [op for op in sm.get_opcodes() if op[0] != "equal"]
    for tag, i1, i2, j1, j2 in reversed(ops):
        _replace_range(p_el, i1, i2, new[j1:j2])


def _segments(p_el):
    segs, pos = [], 0
    for tok in _walk_tokens(p_el):
        if tok[0] == "t":
            n = len(tok[1].text or "")
            segs.append(["t", tok[1], pos, pos + n])
            pos += n
        elif tok[0] == "x":
            segs.append(["x", tok[1], pos, pos + 1])
            pos += 1
        else:
            segs.append(["pic", tok[1], pos, pos])
    return segs


def _replace_range(p_el, i1, i2, ins):
    segs = _segments(p_el)
    tsegs = [s for s in segs if s[0] == "t"]
    # 서식을 물려받을 w:t: 바뀌는 첫 글자 / 삽입이면 바로 앞 글자
    anchor = None
    if i2 > i1:
        anchor = next((s for s in tsegs if s[2] <= i1 < s[3]), None)
    if anchor is None:
        anchor = next((s for s in tsegs if s[2] < i1 <= s[3]), None)
    if anchor is None:
        anchor = next((s for s in tsegs if s[2] >= i1), None) or (tsegs[-1] if tsegs else None)
    if anchor is None:
        _ensure_text_run(p_el)
        return _replace_range(p_el, i1, i2, ins)
    # 1) 지우기 (anchor 이외)
    for s in segs:
        if s is anchor:
            continue
        a, b = max(i1, s[2]), min(i2, s[3])
        if a >= b:
            continue
        if s[0] == "t":
            t = s[1].text or ""
            s[1].text = t[:a - s[2]] + t[b - s[2]:]
            s[1].set(XML_SPACE, "preserve")
        else:
            s[1].getparent().remove(s[1])
    # 2) anchor 안에서 바꾸기
    t = anchor[1].text or ""
    la = min(max(i1 - anchor[2], 0), len(t))
    lb = min(max(i2 - anchor[2], la), len(t))
    head, tail = t[:la], t[lb:]
    if "\t" not in ins and "\n" not in ins:
        anchor[1].text = head + ins + tail
        anchor[1].set(XML_SPACE, "preserve")
        return
    parts = _pieces_to_elems(ins)
    anchor[1].text = head
    anchor[1].set(XML_SPACE, "preserve")
    prev = anchor[1]
    for e in parts + ([_new_t(tail)] if tail else []):
        prev.addnext(e)
        prev = e


def _ensure_text_run(p_el):
    """글자가 하나도 없는 문단에 run 하나 만들기 (문단 기호의 글자 서식 사용)"""
    r = p_el.makeelement(W_R, {})
    ppr = p_el.find(W_PPR)
    if ppr is not None and ppr.find(W_RPR) is not None:
        r.append(copy.deepcopy(ppr.find(W_RPR)))
    r.append(_new_t(""))
    p_el.append(r)


# ====================================================================== 복제 도우미

def _first_text_run(p_el):
    for r in p_el.iter(W_R):
        if r.find(W_T) is not None:
            return r
    return None


def clone_para(template_el, text=""):
    """문단 모양(pPr) + 첫 글자 서식(rPr)만 가져온 새 문단"""
    new = template_el.makeelement(W_P, {})
    ppr = template_el.find(W_PPR)
    if ppr is not None:
        ppr = copy.deepcopy(ppr)
        for tag in ("w:sectPr",):                    # 구역 나누기는 복사하지 않음
            for x in ppr.findall(qn(tag)):
                ppr.remove(x)
        new.append(ppr)
    r = new.makeelement(W_R, {})
    src = _first_text_run(template_el)
    rpr = src.find(W_RPR) if src is not None else (ppr.find(W_RPR) if ppr is not None else None)
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    new.append(r)
    for e in _pieces_to_elems(text) or [_new_t("")]:
        r.append(e)
    return new


def block_el(blk):
    k, obj = blk
    return obj.ref


def insert_after(ref_el, new_el):
    ref_el.addnext(new_el)
    return new_el


def end_anchor(doc, sec):
    """절(하위 포함) 마지막 요소 — 그 뒤에 새 내용을 넣는다"""
    for k, obj in reversed(doc.blocks[sec.block_idx:sec.end_idx]):
        if obj.ref is not None:
            return obj.ref
    return sec.para.ref


def _is_caption(p):
    return bool(RE_CAP_ANY.match(p.orig))


def find_heading_template(rep, doc, level, prefer=None):
    """같은 단계 제목 문단 (없으면 가장 가까운 단계)"""
    if prefer is not None and prefer.para is not None and prefer.level == level:
        return prefer.para.ref
    best = None
    for s in rep.all_sections(doc):
        if s.para is None or s.para.ref is None:
            continue
        if s.level == level:
            return s.para.ref
        if best is None or abs(s.level - level) < abs(best.level - level):
            best = s
    return best.para.ref if best else None


def find_body_template(doc, sec=None):
    """본문 문단(제목·캡션·그림·빈 문단 제외). 같은 절 안을 먼저 찾음"""
    def ok(p):
        return (p.kind == "body" and p.orig.strip() and not p.pic and heading_of(p) is None
                and not _is_caption(p) and not re.match(r"^\s*(\(\d+\)|[①-⑳]|-|·|⦁)", p.orig))
    rng = []
    if sec is not None:
        rng = [b for b in doc.blocks[sec.block_idx:sec.end_idx]]
    for k, obj in rng + doc.blocks:
        if k == "p" and ok(obj):
            return obj.ref
    for k, obj in doc.blocks:
        if k == "p" and obj.kind == "body" and obj.orig.strip():
            return obj.ref
    return None


def find_caption_template(doc, kind):
    """kind: '표' | '그림' — 같은 종류 캡션, 없으면 다른 종류 캡션"""
    other = None
    for p in doc.paras:
        m = RE_CAP_ANY.match(p.orig)
        if m:
            if m.group(1) == kind or (kind == "그림" and m.group(1) == "사진"):
                return p.ref, p.orig
            other = other or (p.ref, p.orig)
    return other or (None, None)


def caption_text_like(sample, kind, chap, no, title):
    """기존 캡션 표기법(<표 5.1 제목>, [그림 5-1] 제목 …)을 따라 새 캡션 글"""
    if sample:
        m = RE_CAP_ANY.match(sample)
        if m:
            sep = sample[m.end(2):m.start(3)].strip() or "."
            pre = sample[:m.start(1)]
            mid = sample[m.end(1):m.start(2)]
            close = {"<": ">", "[": "]", "〈": "〉", "(": ")"}.get(pre.strip()[:1], "")
            rest = sample[m.end(3):]
            if close and close in rest and rest.strip().startswith(close):     # [표 5.1] 제목
                return f"{pre}{kind}{mid}{chap}{sep}{no}{close} {title}".rstrip()
            return f"{pre}{kind}{mid}{chap}{sep}{no} {title}{close}".rstrip()
    return f"<{kind} {chap}.{no} {title}>"


# ====================================================================== 표·그림 만들기

def build_table_like(template_tbl, rows, cols, header=None):
    """기존 표의 테두리·셀 서식·글씨를 본떠 rows×cols 새 표"""
    t = copy.deepcopy(template_tbl)
    trs = [tr for tr in t.findall(W_TR)]
    hdr_tr = copy.deepcopy(trs[0])
    body_tr = copy.deepcopy(trs[1] if len(trs) > 1 else trs[0])
    for tr in trs:
        t.remove(tr)
    grid = t.find(qn("w:tblGrid"))
    total = 0
    if grid is not None:
        total = sum(int(g.get(qn("w:w"), "0") or 0) for g in grid.findall(qn("w:gridCol")))
        for g in list(grid):
            grid.remove(g)
    if not total:
        total = 9000
    cw = int(total / cols)
    if grid is not None:
        for _ in range(cols):
            g = grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(cw)})
            grid.append(g)

    def make_row(proto, texts):
        tr = copy.deepcopy(proto)
        tcs = tr.findall(W_TC)
        cell_proto = copy.deepcopy(tcs[0]) if tcs else None
        # 머리글 행이 아니면 두 번째 칸(값 칸) 모양이 더 일반적
        if len(tcs) > 1 and proto is body_tr:
            cell_proto = copy.deepcopy(tcs[1])
        for tc in tcs:
            tr.remove(tc)
        for ci in range(cols):
            tc = copy.deepcopy(cell_proto)
            tcpr = tc.find(W_TCPR)
            if tcpr is not None:
                for tag in ("w:gridSpan", "w:vMerge", "w:hMerge"):
                    for x in tcpr.findall(qn(tag)):
                        tcpr.remove(x)
                w = tcpr.find(qn("w:tcW"))
                if w is not None:
                    w.set(qn("w:w"), str(cw))
                    w.set(qn("w:type"), "dxa")
            ps = [c for c in tc if c.tag == W_P]
            tmpl = ps[0] if ps else None
            for c in list(tc):
                if c.tag in (W_P, W_TBL):
                    tc.remove(c)
            txt = texts[ci] if texts and ci < len(texts) else ""
            tc.append(clone_para(tmpl, txt) if tmpl is not None else _plain_p(txt))
            tr.append(tc)
        return tr

    t.append(make_row(hdr_tr, header))
    for _ in range(rows - 1):
        t.append(make_row(body_tr, None))
    return t


def _plain_p(text):
    p = _T_PROTO.makeelement(W_P, {})
    r = p.makeelement(W_R, {})
    r.append(_new_t(text))
    p.append(r)
    return p


def simple_table(docx_document, rows, cols, header=None):
    """보고서에 표가 하나도 없을 때"""
    tb = docx_document.add_table(rows=rows, cols=cols)
    try:
        tb.style = "Table Grid"
    except Exception:
        pass
    if header:
        for i, h in enumerate(header[:cols]):
            tb.cell(0, i).text = h
    el = tb._tbl
    el.getparent().remove(el)
    return el


def picture_run(docx_document, image_path, width_cm=None, width_emu=None):
    """python-docx 로 그림 run 을 만들고 떼어 낸다"""
    tmp = docx_document.add_paragraph()
    run = tmp.add_run()
    if width_emu:
        run.add_picture(image_path, width=Emu(width_emu))
    else:
        run.add_picture(image_path, width=Cm(width_cm or 12))
    r = run._r
    tmp._p.remove(r)
    tmp._p.getparent().remove(tmp._p)
    return r


def existing_picture(doc):
    """(그림 문단 요소, 그림 너비 EMU) — 보고서에 이미 있는 그림"""
    for p in doc.paras:
        if p.pic and p.ref is not None:
            ext = p.ref.find(".//" + qn("wp:extent"))
            cx = int(ext.get("cx")) if ext is not None else None
            return p.ref, cx
    return None, None


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def pictures_in(p_el):
    """문단 안 그림들: [{'inline': wp:inline|wp:anchor, 'rid', 'cx', 'cy'}]"""
    out = []
    for d in p_el.iter(W_DRAWING):
        box = d.find(qn("wp:inline"))
        if box is None:
            box = d.find(qn("wp:anchor"))
        if box is None:
            continue
        ext = box.find(qn("wp:extent"))
        blip = box.find(".//{%s}blip" % A_NS)
        out.append({"inline": box, "rid": blip.get(R_EMBED) if blip is not None else None,
                    "cx": int(ext.get("cx")) if ext is not None else 0,
                    "cy": int(ext.get("cy")) if ext is not None else 0})
    return out


def image_blob(docx_document, rid):
    try:
        return docx_document.part.related_parts[rid].blob
    except Exception:
        return None


def set_picture_size(box, cx, cy):
    cx, cy = int(cx), int(cy)
    ext = box.find(qn("wp:extent"))
    if ext is not None:
        ext.set("cx", str(cx))
        ext.set("cy", str(cy))
    for e in box.iter("{%s}ext" % A_NS):
        if e.getparent().tag == "{%s}xfrm" % A_NS:
            e.set("cx", str(cx))
            e.set("cy", str(cy))


def text_area_emu(docx_document):
    """본문 영역 (폭, 높이) EMU"""
    s = docx_document.sections[0]
    return (int(s.page_width - s.left_margin - s.right_margin),
            int(s.page_height - s.top_margin - s.bottom_margin))


def fit_picture_size(image_path, want_cx, area):
    """원본 비율 유지, 본문 폭·높이(85%) 안에 들어가게"""
    try:
        from PIL import Image
        with Image.open(image_path) as im:
            w, h = im.size
    except Exception:
        w, h = 4, 3
    tw, th = area
    cx = min(int(want_cx or tw * 0.8), tw)
    cy = int(cx * h / max(1, w))
    if cy > th * 0.85:
        cy = int(th * 0.85)
        cx = int(cy * w / max(1, h))
    return cx, cy


def picture_para(doc, image_path, width_cm=None):
    tmpl, cx = existing_picture(doc)
    if tmpl is not None:
        new = tmpl.makeelement(W_P, {})
        ppr = tmpl.find(W_PPR)
        if ppr is not None:
            new.append(copy.deepcopy(ppr))
    else:
        new = _T_PROTO.makeelement(W_P, {})
        ppr = new.makeelement(W_PPR, {})
        jc = ppr.makeelement(qn("w:jc"), {qn("w:val"): "center"})
        ppr.append(jc)
        new.append(ppr)
    area = text_area_emu(doc.docx)
    want = int(width_cm * 360000) if width_cm else cx           # 기존 그림 폭 → 없으면 본문 폭의 80%
    wcx, wcy = fit_picture_size(image_path, want, area)
    tmp = doc.docx.add_paragraph()
    run = tmp.add_run()
    run.add_picture(image_path, width=Emu(wcx), height=Emu(wcy))
    r = run._r
    tmp._p.remove(r)
    tmp._p.getparent().remove(tmp._p)
    new.append(r)
    return new


def tidy_header_footer(docx_document):
    """머리말/꼬리말이 표로 끝나 뒤에 붙은 빈 문단 → 1pt (안 그러면 쪽마다 본문 위·아래에 빈 줄)"""
    n = 0
    for sec in docx_document.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer,
                   sec.even_page_header, sec.even_page_footer):
            try:
                if hf.is_linked_to_previous:
                    continue
                el = hf._element
            except Exception:
                continue
            kids = [c for c in el if c.tag in (W_P, W_TBL)]
            if len(kids) < 2 or kids[-1].tag != W_P or kids[-2].tag != W_TBL:
                continue
            last = kids[-1]
            if "".join(t.text or "" for t in last.iter(W_T)).strip() or last.find(".//" + qn("w:drawing")) is not None:
                continue
            ppr = last.find(W_PPR)
            if ppr is None:
                ppr = last.makeelement(W_PPR, {})
                last.insert(0, ppr)
            sp = ppr.find(qn("w:spacing"))
            if sp is not None and sp.get(qn("w:line")) == "20":
                continue
            if sp is None:
                sp = ppr.makeelement(qn("w:spacing"), {})
                ppr.append(sp)
            for k, v in (("w:before", "0"), ("w:after", "0"), ("w:line", "20"), ("w:lineRule", "exact")):
                sp.set(qn(k), v)
            rpr = ppr.find(W_RPR)
            if rpr is None:
                rpr = ppr.makeelement(W_RPR, {})
                ppr.append(rpr)
            sz = rpr.find(qn("w:sz"))
            if sz is None:
                sz = rpr.makeelement(qn("w:sz"), {})
                rpr.append(sz)
            sz.set(qn("w:val"), "2")
            n += 1
    return n


def save_docx(doc, path=None):
    tidy_header_footer(doc.docx)
    doc.docx.save(path or doc.path)


def add_table_row(tbl_el):
    """표 맨 아래에 마지막 행 모양 그대로 빈 행 추가"""
    trs = [tr for tr in tbl_el.findall(W_TR)]
    if not trs:
        return None
    new = copy.deepcopy(trs[-1])
    for tc in new.findall(W_TC):
        tcpr = tc.find(W_TCPR)
        if tcpr is not None:
            for x in tcpr.findall(qn("w:vMerge")):
                tcpr.remove(x)
        ps = [c for c in tc if c.tag == W_P]
        tmpl = ps[0] if ps else None
        for c in list(tc):
            if c.tag in (W_P, W_TBL):
                tc.remove(c)
        tc.append(clone_para(tmpl, "") if tmpl is not None else _plain_p(""))
    trs[-1].addnext(new)
    return new


def add_picture_to_cell(doc, tc, image_path):
    """표 칸 안에 그림 (칸 폭에 맞춤, 비율 유지)"""
    tcpr = tc.find(W_TCPR)
    w_tw = None
    if tcpr is not None and tcpr.find(qn("w:tcW")) is not None:
        try:
            w_tw = int(tcpr.find(qn("w:tcW")).get(qn("w:w"), "0"))
        except ValueError:
            w_tw = None
    area = text_area_emu(doc.docx)
    cell_w = int(w_tw * 635 * 0.92) if w_tw else int(area[0] * 0.4)
    cx, cy = fit_picture_size(image_path, cell_w, (cell_w, int(area[1] * 0.6)))
    tmp = doc.docx.add_paragraph()
    run = tmp.add_run()
    run.add_picture(image_path, width=Emu(cx), height=Emu(cy))
    r = run._r
    tmp._p.remove(r)
    tmp._p.getparent().remove(tmp._p)
    ps = [c for c in tc if c.tag == W_P]
    target = None
    for p in ps:                                      # 빈 문단이 있으면 거기에
        txt, _, pic = para_text(p)
        if not txt.strip() and not pic:
            target = p
            break
    if target is None:
        target = clone_para(ps[0], "") if ps else _plain_p("")
        (ps[-1].addnext(target) if ps else tc.append(target))
    for x in list(target):
        if x.tag == W_R:
            target.remove(x)
    ppr = target.find(W_PPR)
    if ppr is None:
        ppr = target.makeelement(W_PPR, {})
        target.insert(0, ppr)
    sp = ppr.find(qn("w:spacing"))
    if sp is not None:                                # 고정 줄 간격이면 그림이 잘리므로 해제
        for a in (qn("w:line"), qn("w:lineRule")):
            if a in sp.attrib:
                del sp.attrib[a]
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = ppr.makeelement(qn("w:jc"), {})
        ppr.append(jc)
    jc.set(qn("w:val"), "center")
    target.append(r)
    return target


def set_cell_paragraphs(cell_paras_el, texts):
    """칸의 글 문단들(cell_paras_el, 그림 문단 제외)을 texts 개수에 맞게 늘리거나 줄임 (글은 바꾸지 않음)
    → 새 목록 [w:p] (texts 와 같은 길이)"""
    els = list(cell_paras_el)
    while len(els) < len(texts):
        n = clone_para(els[-1], "")
        els[-1].addnext(n)
        els.append(n)
    while len(els) > max(1, len(texts)):
        e = els.pop()
        e.getparent().remove(e)
    return els
