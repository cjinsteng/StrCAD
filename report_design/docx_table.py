# -*- coding: utf-8 -*-
"""
docx_table.py — Word 표 '레이아웃' 기능 (Word 의 [표 레이아웃] 탭과 같은 동작)

행/열 삽입(위·아래·왼쪽·오른쪽), 행/열/표 삭제, 셀 병합·분할, 자동 맞춤(내용/창/고정),
행 높이·열 너비 지정, 같게 맞추기, 셀 맞춤(9방향), 머리글 행 반복.

좌표는 (행 번호, 격자 열 번호). 병합된 셀은 왼쪽 위 칸 좌표로 가리킨다.
"""
import copy

from docx.oxml.ns import qn

from docx_doc import clone_para, _plain_p, W_P, W_TBL, W_TR, W_TC, W_TCPR, W_PPR

TWIPS_PER_CM = 1440 / 2.54


def _get(parent, tag, create=False, first=False):
    el = parent.find(qn(tag))
    if el is None and create:
        el = parent.makeelement(qn(tag), {})
        if first:
            parent.insert(0, el)
        else:
            parent.append(el)
    return el


def _tcpr(tc):
    pr = tc.find(W_TCPR)
    if pr is None:
        pr = tc.makeelement(W_TCPR, {})
        tc.insert(0, pr)
    return pr


def _span(tc):
    pr = tc.find(W_TCPR)
    g = pr.find(qn("w:gridSpan")) if pr is not None else None
    return int(g.get(qn("w:val"), "1")) if g is not None else 1


def _set_span(tc, n):
    pr = _tcpr(tc)
    g = pr.find(qn("w:gridSpan"))
    if n <= 1:
        if g is not None:
            pr.remove(g)
        return
    if g is None:
        g = pr.makeelement(qn("w:gridSpan"), {})
        w = pr.find(qn("w:tcW"))
        (w.addnext(g) if w is not None else pr.insert(0, g))
    g.set(qn("w:val"), str(n))


def _vm(tc):
    pr = tc.find(W_TCPR)
    v = pr.find(qn("w:vMerge")) if pr is not None else None
    if v is None:
        return None
    return v.get(qn("w:val"), "continue")


def _set_vm(tc, val):
    pr = _tcpr(tc)
    v = pr.find(qn("w:vMerge"))
    if val is None:
        if v is not None:
            pr.remove(v)
        return
    if v is None:
        v = pr.makeelement(qn("w:vMerge"), {})
        anchor = pr.find(qn("w:gridSpan"))
        if anchor is None:
            anchor = pr.find(qn("w:tcW"))
        (anchor.addnext(v) if anchor is not None else pr.insert(0, v))
    if val == "restart":
        v.set(qn("w:val"), "restart")
    elif qn("w:val") in v.attrib:
        del v.attrib[qn("w:val")]


def _clear_tc(tc):
    """셀 내용 비우기 (첫 문단 모양은 유지)"""
    ps = [c for c in tc if c.tag == W_P]
    tmpl = ps[0] if ps else None
    for c in list(tc):
        if c.tag in (W_P, W_TBL):
            tc.remove(c)
    tc.append(clone_para(tmpl, "") if tmpl is not None else _plain_p(""))


def _content(tc):
    return [c for c in tc if c.tag in (W_P, W_TBL)]


def _has_text(tc):
    return any(t.text and t.text.strip() for t in tc.iter(qn("w:t")))


class TGrid:
    def __init__(self, tbl):
        self.tbl = tbl
        self.reload()

    # ------------------------------------------------------------ 구조 읽기
    def reload(self):
        tbl = self.tbl
        self.trs = [tr for tr in tbl if tr.tag == W_TR]
        grid = _get(tbl, "w:tblGrid", create=True)
        self.gridcols = [g for g in grid.findall(qn("w:gridCol"))]
        self.at = {}
        self.rows = []
        ncol = 0
        for r, tr in enumerate(self.trs):
            row = []
            c = 0
            trpr = tr.find(qn("w:trPr"))
            gb = trpr.find(qn("w:gridBefore")) if trpr is not None else None
            if gb is not None:
                c = int(gb.get(qn("w:val"), "0"))
            for tc in tr.findall(W_TC):
                e = {"tc": tc, "c": c, "span": _span(tc), "vm": _vm(tc), "r": r}
                row.append(e)
                for k in range(c, c + e["span"]):
                    self.at[(r, k)] = e
                c += e["span"]
            ncol = max(ncol, c)
            self.rows.append(row)
        while len(self.gridcols) < ncol:                # tblGrid 가 모자라면 채움
            g = grid.makeelement(qn("w:gridCol"), {qn("w:w"): "1000"})
            grid.append(g)
            self.gridcols.append(g)
        self.ncol = max(ncol, len(self.gridcols))
        self.nrow = len(self.trs)

    def widths(self):
        return [int(g.get(qn("w:w"), "0") or 0) for g in self.gridcols]

    def set_widths(self, ws):
        for g, w in zip(self.gridcols, ws):
            g.set(qn("w:w"), str(max(100, int(w))))
        self._sync_widths()

    def _sync_widths(self):
        ws = self.widths()
        for row in self.rows:
            for e in row:
                pr = _tcpr(e["tc"])
                w = pr.find(qn("w:tcW"))
                if w is None:
                    w = pr.makeelement(qn("w:tcW"), {})
                    pr.insert(0, w)
                w.set(qn("w:w"), str(sum(ws[e["c"]:e["c"] + e["span"]])))
                w.set(qn("w:type"), "dxa")
        tblpr = _get(self.tbl, "w:tblPr", create=True, first=True)
        tw = _get(tblpr, "w:tblW", create=True)
        if tw.get(qn("w:type")) != "pct":
            tw.set(qn("w:w"), str(sum(ws)))
            tw.set(qn("w:type"), "dxa")

    def clamp(self, r, c):
        return max(0, min(r, self.nrow - 1)), max(0, min(c, self.ncol - 1))

    def origin(self, r, c):
        """(r,c) 칸을 덮고 있는 병합 셀의 왼쪽 위 (행, 열)"""
        r, c = self.clamp(r, c)
        e = self.at.get((r, c))
        if e is None:
            return r, c
        c0 = e["c"]
        rr = r
        while rr > 0 and self.at.get((rr, c0), {}).get("vm") == "continue":
            rr -= 1
        return rr, c0

    def extent(self, r, c):
        """병합 셀의 (행 수, 열 수)"""
        r0, c0 = self.origin(r, c)
        e = self.at.get((r0, c0))
        if e is None:
            return 1, 1
        n = 1
        while (r0 + n, c0) in self.at and self.at[(r0 + n, c0)]["vm"] == "continue" \
                and self.at[(r0 + n, c0)]["c"] == c0:
            n += 1
        return n, e["span"]

    def _row_template(self, r):
        return self.trs[r]

    # ------------------------------------------------------------ 행
    def insert_row(self, r, below=True):
        r = max(0, min(r, self.nrow - 1))
        src = self.trs[r]
        new = copy.deepcopy(src)
        for tc in new.findall(W_TC):
            _clear_tc(tc)
        # 세로 병합 안쪽에 넣으면 병합을 이어 준다
        new_tcs = new.findall(W_TC)
        for e, tc in zip(self.rows[r], new_tcs):
            inside = False
            if below:
                nxt = self.at.get((r + 1, e["c"]))
                inside = nxt is not None and nxt["vm"] == "continue" and nxt["c"] == e["c"]
            else:
                inside = e["vm"] == "continue"
            _set_vm(tc, "continue" if inside else None)
        trpr = new.find(qn("w:trPr"))
        if trpr is not None:
            for x in trpr.findall(qn("w:tblHeader")):
                trpr.remove(x)
        (src.addnext(new) if below else src.addprevious(new))
        self.reload()

    def delete_rows(self, rows):
        for r in sorted(set(rows), reverse=True):
            if r >= len(self.trs):
                continue
            # 이 행에서 시작한 세로 병합은 다음 행이 이어 받음
            for e in self.rows[r]:
                if e["vm"] == "restart":
                    nxt = self.at.get((r + 1, e["c"]))
                    if nxt is not None and nxt["vm"] == "continue":
                        below = self.at.get((r + 2, e["c"]))
                        _set_vm(nxt["tc"], "restart" if (below is not None and below["vm"] == "continue") else None)
                        if _has_text(e["tc"]) and not _has_text(nxt["tc"]):
                            for c in list(nxt["tc"]):
                                if c.tag in (W_P, W_TBL):
                                    nxt["tc"].remove(c)
                            for c in _content(e["tc"]):
                                nxt["tc"].append(c)
            self.tbl.remove(self.trs[r])
            self.reload()
        self._fix_lonely_restart()

    def _fix_lonely_restart(self):
        for r, row in enumerate(self.rows):
            for e in row:
                if e["vm"] == "restart":
                    nxt = self.at.get((r + 1, e["c"]))
                    if nxt is None or nxt["vm"] != "continue":
                        _set_vm(e["tc"], None)
        self.reload()

    # ------------------------------------------------------------ 열
    def insert_col(self, c, right=True, new_cell_rows=None):
        c = max(0, min(c, self.ncol - 1))
        ws = self.widths()
        base = ws[c] if c < len(ws) else 1000
        half = max(300, base // 2)
        pos = c + 1 if right else c
        grid = self.tbl.find(qn("w:tblGrid"))
        g = grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(half)})
        if pos < len(self.gridcols):
            self.gridcols[pos].addprevious(g)
        else:
            grid.append(g)
        self.gridcols[c].set(qn("w:w"), str(max(300, base - half)))
        for r, row in enumerate(self.rows):
            e = self.at.get((r, c))
            if e is None:
                continue
            first, last = e["c"], e["c"] + e["span"] - 1
            boundary = (right and last == c) or (not right and first == c)
            if new_cell_rows is not None:
                boundary = r in new_cell_rows
            if boundary:
                tc = copy.deepcopy(e["tc"])
                _clear_tc(tc)
                _set_span(tc, 1)
                (e["tc"].addnext(tc) if right else e["tc"].addprevious(tc))
            else:
                _set_span(e["tc"], e["span"] + 1)
        self.reload()
        self._sync_widths()

    def delete_cols(self, cols):
        for c in sorted(set(cols), reverse=True):
            if c >= self.ncol:
                continue
            for r in range(self.nrow):
                e = self.at.get((r, c))
                if e is None:
                    continue
                if e["span"] > 1:
                    _set_span(e["tc"], e["span"] - 1)
                else:
                    e["tc"].getparent().remove(e["tc"])
            self.gridcols[c].getparent().remove(self.gridcols[c])
            self.reload()
        for tr in list(self.trs):                       # 칸이 하나도 없는 행 정리
            if not tr.findall(W_TC):
                self.tbl.remove(tr)
        self.reload()
        self._sync_widths()

    # ------------------------------------------------------------ 병합 / 분할
    def _expand_rect(self, r1, c1, r2, c2):
        """병합 셀이 걸치면 사각형을 넓힘"""
        changed = True
        while changed:
            changed = False
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    r0, c0 = self.origin(r, c)
                    nr, nc = self.extent(r, c)
                    a, b, cc, d = min(r1, r0), min(c1, c0), max(r2, r0 + nr - 1), max(c2, c0 + nc - 1)
                    if (a, b, cc, d) != (r1, c1, r2, c2):
                        r1, c1, r2, c2 = a, b, cc, d
                        changed = True
        return r1, c1, r2, c2

    def merge(self, r1, c1, r2, c2):
        r1, c1 = self.clamp(r1, c1)
        r2, c2 = self.clamp(r2, c2)
        r1, c1, r2, c2 = self._expand_rect(min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
        if (r1, c1) == (r2, c2):
            return
        keep = self.at[(r1, c1)]["tc"]
        moved = []
        for r in range(r1, r2 + 1):
            row_es = [e for e in self.rows[r] if c1 <= e["c"] <= c2]
            first = row_es[0]["tc"]
            for e in row_es:
                if e["tc"] is not keep and _has_text(e["tc"]):
                    moved += [c for c in _content(e["tc"]) if c.tag == W_TBL or "".join(
                        t.text or "" for t in c.iter(qn("w:t"))).strip()]
            for e in row_es[1:]:
                e["tc"].getparent().remove(e["tc"])
            _set_span(first, c2 - c1 + 1)
            if r2 > r1:
                _set_vm(first, "restart" if r == r1 else "continue")
            else:
                _set_vm(first, None)
            if r > r1:
                _clear_tc(first)
        for c in moved:                                  # 다른 칸 글은 합친 칸 아래에 이어 붙임
            if not _has_text(keep):
                for x in list(keep):
                    if x.tag in (W_P, W_TBL):
                        keep.remove(x)
            keep.append(c)
        self.reload()
        self._sync_widths()

    def split(self, r, c, nrows, ncols):
        """셀 분할 (Word 와 같이): 병합된 셀은 풀어서 나누고, 모자라는 행/열은 표에 추가"""
        r, c = self.clamp(r, c)
        r0, c0 = self.origin(r, c)
        er, ec = self.extent(r0, c0)
        nrows, ncols = max(1, nrows), max(1, ncols)
        # 1) 세로
        if er >= nrows:
            sizes = [er // nrows + (1 if i < er % nrows else 0) for i in range(nrows)]
            k = 0
            for sz in sizes:
                for j in range(sz):
                    tc = self.at[(r0 + k + j, c0)]["tc"]
                    _set_vm(tc, None if sz == 1 else ("restart" if j == 0 else "continue"))
                k += sz
            self.reload()
            block = er
        else:
            for k in range(er):
                _set_vm(self.at[(r0 + k, c0)]["tc"], None)
            self.reload()
            for _ in range(nrows - er):                 # 행 추가: 다른 칸은 세로 병합으로 이어 붙임
                last = r0 + er - 1
                self.insert_row(last, below=True)
                for e in self.rows[last + 1]:
                    if e["c"] == c0:
                        _set_vm(e["tc"], None)
                        continue
                    o_r, o_c = self.origin(last, e["c"])
                    src = self.at[(o_r, o_c)]
                    if src["vm"] is None:
                        _set_vm(src["tc"], "restart")
                    _set_vm(e["tc"], "continue")
                er += 1
                self.reload()
            block = er
        # 2) 가로 — 블록의 모든 행(세로 병합 이어진 칸 포함)을 같이 나눔
        if ncols == 1:
            self._sync_widths()
            return
        rows = list(range(r0, r0 + block))
        if ec >= ncols:
            parts = [ec // ncols + (1 if i < ec % ncols else 0) for i in range(ncols)]
            for rr in rows:
                e = self.at[(rr, c0)]
                tc, vm = e["tc"], e["vm"]
                _set_span(tc, parts[0])
                prev = tc
                for sp in parts[1:]:
                    n = copy.deepcopy(tc)
                    _clear_tc(n)
                    _set_span(n, sp)
                    _set_vm(n, vm)
                    prev.addnext(n)
                    prev = n
            self.reload()
        else:
            for rr in rows:                              # 먼저 1칸씩으로
                e = self.at[(rr, c0)]
                tc, vm = e["tc"], e["vm"]
                if e["span"] > 1:
                    _set_span(tc, 1)
                    prev = tc
                    for _ in range(e["span"] - 1):
                        n = copy.deepcopy(tc)
                        _clear_tc(n)
                        _set_vm(n, vm)
                        prev.addnext(n)
                        prev = n
            self.reload()
            for _ in range(ncols - ec):                 # 모자라는 열: 이 블록만 새 칸, 다른 행은 넓어짐
                self.insert_col(c0 + ec - 1, right=True, new_cell_rows=set(rows))
                for rr in rows[1:]:                     # 세로 병합이 이어지도록
                    e = self.at[(rr, c0 + ec)]
                    top = self.at[(rows[0], c0 + ec)]
                    _set_vm(e["tc"], _vm(self.at[(rr, c0)]["tc"]))
                    _set_vm(top["tc"], _vm(self.at[(rows[0], c0)]["tc"]))
                ec += 1
                self.reload()
        self._sync_widths()

    # ------------------------------------------------------------ 크기 / 맞춤
    def set_row_height(self, rows, cm, rule="atLeast"):
        for r in rows:
            tr = self.trs[r]
            trpr = tr.find(qn("w:trPr"))
            if trpr is None:
                trpr = tr.makeelement(qn("w:trPr"), {})
                tr.insert(0, trpr)
            h = _get(trpr, "w:trHeight", create=True)
            h.set(qn("w:val"), str(int(cm * TWIPS_PER_CM)))
            h.set(qn("w:hRule"), rule)

    def row_height(self, r):
        trpr = self.trs[r].find(qn("w:trPr"))
        h = trpr.find(qn("w:trHeight")) if trpr is not None else None
        return int(h.get(qn("w:val"), "0")) / TWIPS_PER_CM if h is not None else None

    def set_col_width(self, cols, cm):
        ws = self.widths()
        for c in cols:
            ws[c] = int(cm * TWIPS_PER_CM)
        self.set_widths(ws)

    def equal_rows(self, rows):
        hs = [self.row_height(r) or 0.7 for r in rows]
        self.set_row_height(rows, max(hs))

    def equal_cols(self, cols):
        ws = self.widths()
        tot = sum(ws[c] for c in cols)
        for c in cols:
            ws[c] = tot // len(cols)
        self.set_widths(ws)

    def autofit(self, mode, text_width_twips=None):
        """mode: content(내용에 자동 맞춤) | window(창에 자동 맞춤) | fixed(열 너비 고정)"""
        tblpr = _get(self.tbl, "w:tblPr", create=True, first=True)
        lay = _get(tblpr, "w:tblLayout", create=True)
        tw = _get(tblpr, "w:tblW", create=True)
        if mode == "content":
            lay.set(qn("w:type"), "autofit")
            # 글자 수에 비례해 열 너비 재분배 (Word 가 열 때 다시 맞춤)
            lens = [1] * self.ncol
            for (r, c), e in self.at.items():
                if e["span"] == 1 and e["c"] == c:
                    n = max((len("".join(t.text or "" for t in p.iter(qn("w:t")))) for p in e["tc"].iter(W_P)), default=1)
                    lens[c] = max(lens[c], min(40, n))
            tot = sum(self.widths())
            mn = 400
            raw = [max(mn, l * 220) for l in lens]
            scale = min(1.0, tot / sum(raw)) if sum(raw) else 1
            self.set_widths([w * scale for w in raw])
            tw.set(qn("w:type"), "auto")
            tw.set(qn("w:w"), "0")
        elif mode == "window":
            lay.set(qn("w:type"), "fixed")
            ws = self.widths()
            tot = sum(ws) or 1
            target = text_width_twips or tot
            self.set_widths([w * target / tot for w in ws])
        else:
            lay.set(qn("w:type"), "fixed")
            self._sync_widths()

    def align(self, cells, vert, horiz):
        """vert: top|center|bottom, horiz: left|center|right"""
        for (r, c) in cells:
            e = self.at.get((r, c))
            if e is None:
                continue
            tc = e["tc"]
            pr = _tcpr(tc)
            va = _get(pr, "w:vAlign", create=True)
            va.set(qn("w:val"), vert)
            for p in tc.iter(W_P):
                ppr = p.find(W_PPR)
                if ppr is None:
                    ppr = p.makeelement(W_PPR, {})
                    p.insert(0, ppr)
                jc = _get(ppr, "w:jc", create=True)
                jc.set(qn("w:val"), horiz)

    def header_repeat(self, rows, on=True):
        for r in range(self.nrow):
            tr = self.trs[r]
            trpr = tr.find(qn("w:trPr"))
            if trpr is None:
                trpr = tr.makeelement(qn("w:trPr"), {})
                tr.insert(0, trpr)
            h = trpr.find(qn("w:tblHeader"))
            want = on and r <= max(rows)
            if want and h is None:
                trpr.append(trpr.makeelement(qn("w:tblHeader"), {}))
            elif not want and h is not None:
                trpr.remove(h)

    def is_header(self, r):
        trpr = self.trs[r].find(qn("w:trPr"))
        return trpr is not None and trpr.find(qn("w:tblHeader")) is not None


def section_text_width(docx_document):
    s = docx_document.sections[0]
    return int((s.page_width - s.left_margin - s.right_margin) / 635)       # EMU → twips
