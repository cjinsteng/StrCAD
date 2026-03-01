"""
Excel 리포트 생성 모듈
"""

from datetime import datetime
from typing import List, Dict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os
import config


class ExcelExporter:
    """분석 결과를 Excel 파일로 출력하는 클래스"""
    
    def __init__(self):
        """초기화"""
        self.output_folder = config.OUTPUT['folder']
        
        # 출력 폴더 생성
        if not os.path.exists(self.output_folder):
            os.makedirs(self.output_folder)
    
    def export_beam_spans(self, span_data: List[Dict], filename: str = None) -> str:
        """
        보 스팬 데이터를 Excel로 출력
        
        Args:
            span_data: 보 스팬 정보 리스트
            filename: 파일명 (None이면 자동 생성)
        
        Returns:
            생성된 파일 경로
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{config.OUTPUT['span_prefix']}_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_folder, filename)
        
        # 워크북 생성
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "보 스팬 정리"
        
        # 헤더 작성
        headers = ['번호', '보 ID', '레이어', '스팬(mm)', '방향', '시작 기둥', '끝 기둥', '시작 X', '시작 Y', '끝 X', '끝 Y']
        self._write_header(ws, headers)
        
        # 데이터 작성
        decimals = config.ANALYSIS_OPTIONS['report_decimals']
        
        for i, span in enumerate(span_data, start=2):
            ws[f'A{i}'] = i - 1
            ws[f'B{i}'] = span['beam_id']
            ws[f'C{i}'] = span['beam_layer']
            ws[f'D{i}'] = round(span['span_length'], decimals)
            ws[f'E{i}'] = span['beam_direction']
            ws[f'F{i}'] = span['start_column'] if span['start_column'] else '-'
            ws[f'G{i}'] = span['end_column'] if span['end_column'] else '-'
            ws[f'H{i}'] = round(span['start_point'][0], decimals)
            ws[f'I{i}'] = round(span['start_point'][1], decimals)
            ws[f'J{i}'] = round(span['end_point'][0], decimals)
            ws[f'K{i}'] = round(span['end_point'][1], decimals)
        
        # 열 너비 조정
        self._auto_adjust_columns(ws)
        
        # 저장
        wb.save(filepath)
        print(f"✓ 보 스팬 리포트 저장: {filepath}")
        return filepath
    
    def export_discontinuities(self, discontinuity_data: List[Dict], filename: str = None) -> str:
        """
        불연속 구간 데이터를 Excel로 출력
        
        Args:
            discontinuity_data: 불연속 구간 정보 리스트
            filename: 파일명 (None이면 자동 생성)
        
        Returns:
            생성된 파일 경로
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{config.OUTPUT['discontinuity_prefix']}_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_folder, filename)
        
        # 워크북 생성
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "불연속 구간"
        
        # 헤더 작성
        headers = ['번호', '부재 ID', '부재 타입', '레이어', '끝점 타입', 'X 좌표', 'Y 좌표']
        self._write_header(ws, headers)
        
        # 데이터 작성
        decimals = config.ANALYSIS_OPTIONS['report_decimals']
        
        for i, disc in enumerate(discontinuity_data, start=2):
            ws[f'A{i}'] = i - 1
            ws[f'B{i}'] = disc['entity_id']
            ws[f'C{i}'] = disc['entity_type']
            ws[f'D{i}'] = disc['layer']
            ws[f'E{i}'] = disc['end_type']
            ws[f'F{i}'] = round(disc['x'], decimals)
            ws[f'G{i}'] = round(disc['y'], decimals)
        
        # 열 너비 조정
        self._auto_adjust_columns(ws)
        
        # 저장
        wb.save(filepath)
        print(f"✓ 불연속 구간 리포트 저장: {filepath}")
        return filepath
    
    def export_column_spacing(self, spacing_data: Dict, filename: str = None) -> str:
        """
        기둥 간격 데이터를 Excel로 출력
        
        Args:
            spacing_data: 기둥 간격 정보 (X축, Y축)
            filename: 파일명 (None이면 자동 생성)
        
        Returns:
            생성된 파일 경로
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{config.OUTPUT['column_spacing_prefix']}_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_folder, filename)
        
        # 워크북 생성
        wb = openpyxl.Workbook()
        
        # X축 간격 시트
        ws_x = wb.active
        ws_x.title = "X축 간격"
        self._write_spacing_sheet(ws_x, spacing_data['x_spacings'], 'X')
        
        # Y축 간격 시트
        ws_y = wb.create_sheet("Y축 간격")
        self._write_spacing_sheet(ws_y, spacing_data['y_spacings'], 'Y')
        
        # 저장
        wb.save(filepath)
        print(f"✓ 기둥 간격 리포트 저장: {filepath}")
        return filepath
    
    def _write_spacing_sheet(self, ws, spacings: List[Dict], axis: str):
        """기둥 간격 시트 작성"""
        headers = ['번호', f'{axis}축 구간', '시작 좌표', '끝 좌표', '간격(mm)']
        self._write_header(ws, headers)
        
        decimals = config.ANALYSIS_OPTIONS['report_decimals']
        
        for i, spacing in enumerate(spacings, start=2):
            ws[f'A{i}'] = i - 1
            ws[f'B{i}'] = f"{spacing['grid_from']} ~ {spacing['grid_to']}"
            ws[f'C{i}'] = round(spacing['coordinate_from'], decimals)
            ws[f'D{i}'] = round(spacing['coordinate_to'], decimals)
            ws[f'E{i}'] = round(spacing['spacing'], decimals)
        
        self._auto_adjust_columns(ws)
    
    def export_summary(self, summary_data: Dict, filename: str = None) -> str:
        """
        전체 분석 요약을 Excel로 출력
        
        Args:
            summary_data: 요약 정보
            filename: 파일명 (None이면 자동 생성)
        
        Returns:
            생성된 파일 경로
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"분석요약_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_folder, filename)
        
        # 워크북 생성
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "분석 요약"
        
        # 제목
        ws['A1'] = "구조도면 분석 요약"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:B1')
        
        row = 3
        
        # 기둥 정보
        if 'column_summary' in summary_data:
            col_sum = summary_data['column_summary']
            ws[f'A{row}'] = "■ 기둥 정보"
            ws[f'A{row}'].font = Font(bold=True, size=12)
            row += 1
            
            ws[f'A{row}'] = "총 기둥 개수"
            ws[f'B{row}'] = col_sum['total_columns']
            row += 1
            
            ws[f'A{row}'] = "X축 그리드라인"
            ws[f'B{row}'] = col_sum['x_gridlines']
            row += 1
            
            ws[f'A{row}'] = "Y축 그리드라인"
            ws[f'B{row}'] = col_sum['y_gridlines']
            row += 2
        
        # 보 스팬 정보
        if 'beam_summary' in summary_data:
            beam_sum = summary_data['beam_summary']
            ws[f'A{row}'] = "■ 보 스팬 정보"
            ws[f'A{row}'].font = Font(bold=True, size=12)
            row += 1
            
            ws[f'A{row}'] = "총 스팬 개수"
            ws[f'B{row}'] = beam_sum['total_spans']
            row += 1
            
            ws[f'A{row}'] = "평균 스팬 길이(mm)"
            ws[f'B{row}'] = round(beam_sum['avg_span_length'], 0)
            row += 1
            
            ws[f'A{row}'] = "최대 스팬 길이(mm)"
            ws[f'B{row}'] = round(beam_sum['max_span_length'], 0)
            row += 1
            
            ws[f'A{row}'] = "최소 스팬 길이(mm)"
            ws[f'B{row}'] = round(beam_sum['min_span_length'], 0)
            row += 2
        
        # 불연속 구간 정보
        if 'discontinuity_summary' in summary_data:
            disc_sum = summary_data['discontinuity_summary']
            ws[f'A{row}'] = "■ 불연속 구간 정보"
            ws[f'A{row}'].font = Font(bold=True, size=12)
            row += 1
            
            ws[f'A{row}'] = "총 불연속 구간"
            ws[f'B{row}'] = disc_sum['total_discontinuities']
            row += 1
            
            if 'discontinuities_by_layer' in disc_sum:
                for layer, count in disc_sum['discontinuities_by_layer'].items():
                    ws[f'A{row}'] = f"  - {layer}"
                    ws[f'B{row}'] = count
                    row += 1
        
        # 열 너비 조정
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
        
        # 저장
        wb.save(filepath)
        print(f"✓ 분석 요약 리포트 저장: {filepath}")
        return filepath
    
    def _write_header(self, ws, headers: List[str]):
        """헤더 행 작성 및 스타일 적용"""
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
    
    def _auto_adjust_columns(self, ws):
        """열 너비 자동 조정"""
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
